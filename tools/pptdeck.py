"""Course-slide image decks from .pptx files (Practical 3 onward).
- Labels are PowerPoint text boxes/callouts: q-version replaces each label's text with its number
  (box, leader line, highlight + orientation tags untouched); a-version = original labels.
- Notes/title/credits are removed from the image and shown in the reveal instead.
- Rendered via LibreOffice → PDF (no image downsampling) → page raster at the photo's native dpi, cropped."""
import os,re,io,copy,hashlib,subprocess
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE as T
from PIL import Image
import pymupdf as fz
EMU=914400
ORIENT={'cranial','caudal','dorsal','ventral','rostral','medial','lateral','proximal','distal','left','right','apex','base'}
def walk(shapes):
    for sh in shapes:
        if sh.shape_type==T.GROUP: yield from walk(sh.shapes)
        else: yield sh
def is_title(sh):
    try: return sh.is_placeholder and 'TITLE' in str(sh.placeholder_format.type)
    except Exception: return False
def pic_px(sh):
    try:
        im=Image.open(io.BytesIO(sh.image.blob)); return im.size,hashlib.md5(sh.image.blob).hexdigest()[:10]
    except Exception: return None,None
def kind(sh,sw):
    t=sh.text_frame.text.strip()
    if is_title(sh): return 'title'
    w=t.split()
    if re.match(r'^(note|this dissection|this image|\*\*)',t.lower()): return 'note'
    if sw>11*EMU and (sh.left or 0)<0.3*sw and len(w)>6: return 'note'
    lines=[l.strip() for l in t.split('\n') if l.strip()]
    if lines and all(l.lower().strip('.:()') in ORIENT for l in lines): return 'orient'
    if lines and len(lines[0].split())<=7: return 'label'
    return 'desc'
def groups(tf):
    """split a text frame into label groups at blank paragraphs → [[para_idx,...],...]"""
    gs=[];cur=[]
    for i,p in enumerate(tf.paragraphs):
        if p.text.strip(): cur.append(i)
        elif cur: gs.append(cur); cur=[]
    if cur: gs.append(cur)
    return gs
def analyze(path):
    prs=Presentation(path); sw,sh_=prs.slide_width,prs.slide_height; out=[]
    for si,s in enumerate(prs.slides):
        d={'i':si,'title':'','notes':[],'labels':[],'orient':[],'pics':[],'desc':[],'hl':0,'sw':sw,'sh':sh_}
        for sh in walk(s.shapes):
            box=[(sh.left or 0)/EMU,(sh.top or 0)/EMU,((sh.left or 0)+(sh.width or 0))/EMU,((sh.top or 0)+(sh.height or 0))/EMU]
            if sh.shape_type==T.PICTURE or (getattr(sh,'is_placeholder',False) and hasattr(sh,'image')):
                px,hsh=pic_px(sh)
                if px and (box[2]-box[0])*(box[3]-box[1])>1.5: d['pics'].append({'box':box,'px':px,'hash':hsh})
                continue
            if sh.has_text_frame and sh.text_frame.text.strip():
                k=kind(sh,sw); t=sh.text_frame.text.strip()
                if k=='title': d['title']=re.sub(r'\s+',' ',t)
                elif k=='note': d['notes'].append(t)
                elif k=='orient': d['orient'].append(box)
                elif k=='desc': d['desc'].append(t)
                else:
                    for g in groups(sh.text_frame):
                        ps=[sh.text_frame.paragraphs[i].text.strip() for i in g]
                        d['labels'].append({'name':ps[0],'more':' '.join(ps[1:]),'box':box,'shape':sh.shape_id,'paras':g})
            elif sh.shape_type in (T.FREEFORM,T.AUTO_SHAPE,T.LINE) or 'Connector' in sh.name: d['hl']+=1
        out.append(d)
    return out
def make_versions(path,outdir,numbering):
    """numbering: {slide_idx: {(shape_id,first_para): n}}; writes q.pptx + a.pptx"""
    for ver in ('q','a'):
        prs=Presentation(path)
        for si,s in enumerate(prs.slides):
            for sh in list(walk(s.shapes)):
                if not sh.has_text_frame or not sh.text_frame.text.strip(): continue
                k=kind(sh,prs.slide_width)
                if k in('title','note') or (k=='desc' and ver=='q'):
                    sh._element.getparent().remove(sh._element); continue
                if k=='label' and ver=='q':
                    nums=numbering.get(si,{})
                    for g in groups(sh.text_frame):
                        n=nums.get((sh.shape_id,g[0]))
                        for j,pi in enumerate(g):
                            p=sh.text_frame.paragraphs[pi]; runs=p.runs
                            for r in runs[1:]: r._r.getparent().remove(r._r)
                            if runs:
                                runs[0].text=(str(n) if (j==0 and n is not None) else ('' if j else '•'))
                                runs[0].font.bold=True
                                if j==0 and n is not None:
                                    from pptx.util import Pt
                                    if (runs[0].font.size or Pt(12))<Pt(16): runs[0].font.size=Pt(16)
        prs.save(f'{outdir}/{ver}.pptx')
        subprocess.run(['python3','/mnt/skills/public/pptx/scripts/office/soffice.py','--headless','--convert-to','pdf:impress_pdf_Export:{"ReduceImageResolution":{"type":"boolean","value":"false"},"Quality":{"type":"long","value":"95"}}','--outdir',outdir,f'{outdir}/{ver}.pptx'],capture_output=True,timeout=280)
def crop_render(pdf,si,d,out,pad=0.05):
    pics=d['pics']; 
    if not pics: return None
    x0=min(p['box'][0] for p in pics); y0=min(p['box'][1] for p in pics); x1=max(p['box'][2] for p in pics); y1=max(p['box'][3] for p in pics)
    for b in [l['box'] for l in d['labels']]+[[o[0]-0.1,o[1],o[2]+0.35,o[3]] for o in d['orient']]:
        x0=min(x0,b[0]);y0=min(y0,b[1]);x1=max(x1,b[2]);y1=max(y1,b[3])
    W,H=d['sw']/EMU,d['sh']/EMU; x0=max(0,x0-pad);y0=max(0,y0-pad);x1=min(W,x1+pad);y1=min(H,y1+pad)
    dpi=max(p['px'][0]/(p['box'][2]-p['box'][0]) for p in pics); dpi=int(min(300,max(150,dpi)))
    pg=fz.open(pdf)[si]; S=pg.rect.width/W
    pm=pg.get_pixmap(clip=fz.Rect(x0*S,y0*S,x1*S,y1*S),dpi=dpi)
    Image.frombytes('RGB',(pm.width,pm.height),pm.samples).save(out,quality=88)
    return out
